import io
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader
from docx import Document


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extrai texto de arquivo PDF usando pypdf."""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        extracted_text = []
        for index, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                extracted_text.append(page_text.strip())

        full_text = "\n\n".join(extracted_text).strip()
        if not full_text:
            raise HTTPException(
                status_code=400,
                detail="O arquivo PDF não contém texto legível (pode ser uma imagem escaneada).",
            )
        return full_text
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Erro ao processar arquivo PDF: {str(e)}",
        )


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extrai texto de arquivo DOCX usando python-docx."""
    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Também extrair texto dentro de tabelas se houver
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text and cell_text not in paragraphs:
                        paragraphs.append(cell_text)

        full_text = "\n\n".join(paragraphs).strip()
        if not full_text:
            raise HTTPException(
                status_code=400,
                detail="O arquivo DOCX está vazio ou sem texto legível.",
            )
        return full_text
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Erro ao processar arquivo DOCX: {str(e)}",
        )


async def extract_resume_text(file: UploadFile) -> str:
    """Valida formato e extrai texto de arquivo PDF ou DOCX."""
    filename = file.filename or ""
    extension = filename.lower().split(".")[-1] if "." in filename else ""

    if extension not in ["pdf", "docx"]:
        raise HTTPException(
            status_code=400,
            detail="Formato de arquivo inválido. Apenas arquivos PDF (.pdf) e Word (.docx) são aceitos.",
        )

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="Arquivo enviado está vazio.",
        )

    if extension == "pdf":
        return extract_text_from_pdf(file_bytes)
    elif extension == "docx":
        return extract_text_from_docx(file_bytes)

    raise HTTPException(status_code=400, detail="Formato não suportado.")
